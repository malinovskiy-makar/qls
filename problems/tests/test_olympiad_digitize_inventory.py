"""Команда `olympiad_digitize` (inventory, render) — опись и отрисовка.

Папка аудита выдуманная, но устроена как настоящая папка МОШ
(`weconomics-data\\olympiads\\audit_mosh_20261009`): `reference_events.jsonl`
с `task_files`/`solution_files`, `raw\\pdf`, `raw\\news_files` с побайтной
копией, `raw\\probe`, html. PDF рисуются PyMuPDF прямо в тесте; кириллица —
встроенным шрифтом `china-s` (в `helv` её нет, текст слоя был бы точками).
База не нужна: тесты — `SimpleTestCase`.
"""
import json
import os
import shutil
import tempfile
from io import StringIO
from pathlib import Path

import fitz
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase

from problems.olympiad_audit import digitize, registry

AUDIT = registry.get('mosh').audit_dir

LONG = ('Задача 1. Фирма-монополист продаёт товар на двух рынках с разной '
        'эластичностью спроса. Найдите цены на каждом рынке и объясните, '
        'почему ценовая дискриминация выгодна фирме. ') * 3


def make_pdf(path, lines=(), pages=1, size=None, rect_only=False):
    doc = fitz.open()
    for _ in range(pages):
        page = (doc.new_page(width=size[0], height=size[1]) if size
                else doc.new_page())
        if rect_only:  # «скан»: только картинка-прямоугольник, слоя нет
            page.draw_rect(fitz.Rect(50, 50, 300, 300), color=(0, 0, 0),
                           fill=(0.5, 0.5, 0.5))
        y = 60
        for line in lines:
            page.insert_textbox(fitz.Rect(40, y, 560, y + 200), line,
                                fontname='china-s', fontsize=9)
            y += 210
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def make_docx(path):
    import docx
    document = docx.Document()
    document.add_paragraph('Задание 1. Найдите равновесную цену.')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Цена', 'Объём'
    table.cell(1, 0).text, table.cell(1, 1).text = '10', '20'
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


class DigitizeTestBase(SimpleTestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix='digitize_'))
        self.addCleanup(shutil.rmtree, self.root, ignore_errors=True)
        self.audit = self.root / AUDIT
        raw = self.audit / 'raw'
        pdf = raw / 'pdf'
        make_pdf(pdf / 'econ-t9.pdf', [LONG])
        make_pdf(pdf / 'econ-s9.pdf', ['Решения. ' + LONG])
        make_pdf(pdf / 'econ-both10.pdf', ['Условия и решения. ' + LONG])
        make_pdf(pdf / 'scan-t11.pdf', rect_only=True, pages=2)
        make_pdf(pdf / 'ans-econ-8-final.pdf', ['Ответы. ' + LONG])
        make_pdf(pdf / 'kriterii_8.pdf', ['Критерии проверки. ' + LONG])
        make_pdf(pdf / 'tasks-econ-7-final.pdf', [LONG])
        make_pdf(pdf / 'tasks_and_solutions_7.pdf', [LONG])
        make_pdf(pdf / 'protokol_rezultatov_2020.pdf', ['Иванов Иван 10 баллов'])
        make_pdf(pdf / 'itogi.pdf', ['Список победителей и призёров\n'
                                     'Петров Пётр Петрович — 95'])
        make_pdf(pdf / 'big.pdf', [LONG], size=(2000, 1000))
        make_docx(pdf / 'dist-9-1.docx')
        news = raw / 'news_files'
        news.mkdir(parents=True)
        shutil.copyfile(pdf / 'econ-t9.pdf', news / 'copy-of-t9.pdf')
        make_pdf(raw / 'probe' / 'tasks-fg-5-6-otbor.pdf', [LONG])
        (raw / 'html').mkdir()
        (raw / 'html' / 'page.html').write_text('<html>протокол</html>',
                                                encoding='utf-8')
        (raw / 'fetch_log.jsonl').write_text('', encoding='utf-8')

        def event(event_id, tasks=(), solutions=()):
            return {'event_id': event_id, 'task_files': [{'file': f} for f in tasks],
                    'solution_files': [{'file': f} for f in solutions]}
        events = [
            event('mosh-2020-final-9-v1', ['econ-t9.pdf'], ['econ-s9.pdf']),
            event('mosh-2020-final-10-v1', ['econ-both10.pdf'], ['econ-both10.pdf']),
            event('mosh-2020-final-11-v1', ['scan-t11.pdf']),
            event('mosh-2020-qualifying-9-v1', ['dist-9-1.docx']),
            event('mosh-2020-final-12-v1', ['big.pdf']),
        ]
        with open(self.audit / 'reference_events.jsonl', 'w', encoding='utf-8') as fh:
            for e in events:
                fh.write(json.dumps(e, ensure_ascii=False) + '\n')

    def run_cmd(self, *args):
        out = StringIO()
        call_command('olympiad_digitize', *args, '--data-root', str(self.root),
                     stdout=out)
        return out.getvalue()

    def inventory(self):
        self.run_cmd('inventory', '--olympiad', 'mosh')
        rows = digitize.read_jsonl(self.audit / 'digitized' / 'inventory.jsonl')
        return {r['file'].split('/', 1)[1]: r for r in rows}


class InventoryTests(DigitizeTestBase):
    def test_olympiad_required(self):
        for sub in ('inventory', 'render'):
            with self.assertRaisesMessage(CommandError, '--olympiad'):
                self.run_cmd(sub)
        self.assertFalse((self.audit / 'digitized').exists())

    def test_kind_from_reference_and_name(self):
        rows = self.inventory()
        expected = {
            'pdf/econ-t9.pdf': ('tasks', 'reference_events'),
            'pdf/econ-s9.pdf': ('solutions', 'reference_events'),
            'pdf/econ-both10.pdf': ('mixed', 'reference_events'),
            'pdf/ans-econ-8-final.pdf': ('solutions', 'name'),
            'pdf/kriterii_8.pdf': ('criteria', 'name'),
            'pdf/tasks-econ-7-final.pdf': ('tasks', 'name'),
            'pdf/tasks_and_solutions_7.pdf': ('mixed', 'name'),
        }
        for name, (kind, source) in expected.items():
            self.assertEqual((rows[name]['kind'], rows[name]['kind_source']),
                             (kind, source), name)
            self.assertTrue(rows[name]['digitize'], name)
        self.assertEqual(rows['pdf/econ-t9.pdf']['event_ids'], ['mosh-2020-final-9-v1'])
        self.assertTrue(rows['pdf/econ-t9.pdf']['text_layer'])
        scan = rows['pdf/scan-t11.pdf']
        self.assertEqual((scan['text_layer'], scan['pages_with_text'], scan['pages']),
                         (False, 0, 2))
        self.assertTrue(scan['digitize'])
        self.assertEqual(rows['html/page.html']['exclude_reason'], digitize.REASON_NOT_DOC)
        self.assertEqual(rows['probe/tasks-fg-5-6-otbor.pdf']['exclude_reason'],
                         digitize.REASON_OTHER_OLYMPIAD_MOSH)

    def test_protocols_excluded(self):
        out = self.run_cmd('inventory', '--olympiad', 'mosh')
        rows = {r['file']: r for r in digitize.read_jsonl(
            self.audit / 'digitized' / 'inventory.jsonl')}
        for name in ('raw/pdf/protokol_rezultatov_2020.pdf', 'raw/pdf/itogi.pdf'):
            self.assertFalse(rows[name]['digitize'], name)
            self.assertEqual(rows[name]['exclude_reason'], digitize.REASON_PERSONAL)
            self.assertIn(name, out)
        self.assertIn('первая страница', rows['raw/pdf/itogi.pdf']['personal_data_hit'])
        events = digitize.read_jsonl(self.audit / 'digitized' / 'events_files.jsonl')
        listed = {f['file'] for e in events for key in
                  ('task_files', 'solution_files', 'criteria_files') for f in e[key]}
        self.assertNotIn('raw/pdf/itogi.pdf', listed)
        self.assertNotIn('raw/pdf/protokol_rezultatov_2020.pdf', listed)
        self.assertIn('raw/pdf/econ-t9.pdf', listed)

    def test_duplicate_marked(self):
        rows = self.inventory()
        dup = rows['news_files/copy-of-t9.pdf']
        self.assertEqual(dup['duplicate_of'], 'raw/pdf/econ-t9.pdf')
        self.assertFalse(dup['digitize'])
        self.assertTrue(dup['exclude_reason'].startswith('дубль'))
        self.assertEqual(dup['page_dir'], rows['pdf/econ-t9.pdf']['page_dir'])
        self.assertEqual(dup['event_ids'], ['mosh-2020-final-9-v1'])
        self.assertEqual(dup.get('edition'), 'news')

    def test_docx_text_directly(self):
        rows = self.inventory()
        row = rows['pdf/dist-9-1.docx']
        self.assertEqual((row['kind'], row['digitize'], row['exclude_reason']),
                         ('tasks', False, digitize.REASON_DOCX))
        md = (self.audit / row['docx_md']).read_text(encoding='utf-8')
        self.assertIn('Найдите равновесную цену', md)
        self.assertIn('| Цена | Объём |', md)

    def test_sum_by_category_equals_files(self):
        out = self.run_cmd('inventory', '--olympiad', 'mosh')
        rows = digitize.read_jsonl(self.audit / 'digitized' / 'inventory.jsonl')
        files = [p for p in (self.audit / 'raw').rglob('*') if p.is_file()]
        self.assertEqual(len(rows), len(files))
        summary = digitize.summarize(rows)
        self.assertEqual(sum(summary['by_kind'].values()), len(files))
        self.assertEqual(sum(summary['by_reason'].values()), len(files))
        self.assertIn('сходится', out)

    def test_events_files_lists(self):
        self.inventory()
        events = {e['event_id']: e for e in digitize.read_jsonl(
            self.audit / 'digitized' / 'events_files.jsonl')}
        both = events['mosh-2020-final-10-v1']
        self.assertEqual([f['file'] for f in both['task_files']],
                         ['raw/pdf/econ-both10.pdf'])
        self.assertEqual([f['file'] for f in both['solution_files']],
                         ['raw/pdf/econ-both10.pdf'])
        self.assertFalse(both['new'])
        self.assertEqual([f['file'] for f in
                          events['mosh-2020-qualifying-9-v1']['task_files']],
                         ['raw/pdf/dist-9-1.docx'])


class RenderTests(DigitizeTestBase):
    def test_render_png_txt_once(self):
        rows = self.inventory()
        out = self.run_cmd('render', '--olympiad', 'mosh')
        digitized = self.audit / 'digitized'
        manifest = digitize.read_jsonl(digitized / 'pages_manifest.jsonl')
        targets = [r for r in rows.values() if r['digitize'] and not r['duplicate_of']]
        pages = sum(r['pages'] for r in targets)
        self.assertEqual(len(manifest), pages)
        self.assertIn(f'Отрисовано страниц: {pages};', out)
        t9 = digitized / rows['pdf/econ-t9.pdf']['page_dir']
        self.assertTrue((t9 / 'p1.png').is_file())
        self.assertIn('Фирма-монополист', (t9 / 'p1.txt').read_text(encoding='utf-8'))
        scan = digitized / rows['pdf/scan-t11.pdf']['page_dir']
        self.assertEqual((scan / 'p2.txt').read_text(encoding='utf-8'), '')
        self.assertTrue((scan / 'p2.png').is_file())
        # дубль и протоколы не рисуются
        self.assertNotIn('raw/news_files/copy-of-t9.pdf', {m['file'] for m in manifest})
        self.assertNotIn('raw/pdf/itogi.pdf', {m['file'] for m in manifest})
        page_dirs = {p.name for p in (digitized / 'pages').iterdir()}
        self.assertEqual(len(page_dirs), len(targets))

        mtimes = {p: p.stat().st_mtime_ns for p in (digitized / 'pages').rglob('*')}
        again = self.run_cmd('render', '--olympiad', 'mosh')
        self.assertIn('Отрисовано страниц: 0;', again)
        self.assertEqual(
            {p: p.stat().st_mtime_ns for p in (digitized / 'pages').rglob('*')}, mtimes)
        self.assertEqual(len(digitize.read_jsonl(digitized / 'pages_manifest.jsonl')),
                         pages)

    def test_long_side_capped(self):
        rows = self.inventory()
        self.run_cmd('render', '--olympiad', 'mosh')
        manifest = digitize.read_jsonl(self.audit / 'digitized' / 'pages_manifest.jsonl')
        big = [m for m in manifest if m['file'] == 'raw/pdf/big.pdf'][0]
        self.assertEqual(max(big['width'], big['height']), digitize.MAX_SIDE_PX)
        self.assertEqual(big['scaled_to_max_side'], digitize.MAX_SIDE_PX)
        normal = [m for m in manifest if m['file'] == 'raw/pdf/econ-t9.pdf'][0]
        self.assertNotIn('scaled_to_max_side', normal)
        self.assertEqual(rows['pdf/econ-t9.pdf']['pages'], 1)
        self.assertAlmostEqual(normal['width'], 595 * 200 / 72, delta=2)  # A4, 200 dpi

    def test_limit(self):
        self.inventory()
        out = self.run_cmd('render', '--olympiad', 'mosh', '--limit', '1')
        self.assertEqual(len(os.listdir(self.audit / 'digitized' / 'pages')), 1)
        self.assertIn('Отрисовано страниц: 1;', out)


class HelperTests(SimpleTestCase):
    def test_kind_from_name(self):
        cases = {
            'tasks-9-econ-otbor-2008-9.pdf': 'tasks',
            'ans-9-econ-otbor-2008-9.pdf': 'solutions',
            'MOSH-2018-8_resheniya_i_kriterii.pdf': 'solutions',
            'krit2019-10.pdf': 'criteria',
            'Kriterii_MOSH_8-9_finalnyie.pdf': 'criteria',
            '2-Klyuchi_MOSH_2020.pdf': 'solutions',
            '7_klass_2026_s_kriteriyami.pdf': 'mixed',
            '2012_tasks_and_solutions_10.pdf': 'mixed',
            '2021_schedule_final_x.pdf': 'other',
            '2019_5.pdf': '',
        }
        for name, kind in cases.items():
            self.assertEqual(digitize.kind_from_name(name), kind, name)

    def test_personal_data_not_triggered_by_tasks(self):
        # «результат» и «победитель» в условии — не протокол
        self.assertEqual(digitize.personal_data_hit(
            'tasks-econ-9.pdf', 'В результате аукциона победитель платит цену'), '')
        self.assertTrue(digitize.personal_data_hit('x.pdf', 'Протокол результатов олимпиады'))

    def test_default_data_root(self):
        self.assertEqual(digitize.default_data_root('/a/qls', {'OLYMPIAD_DATA_ROOT': '/d'}),
                         Path('/d'))
        self.assertEqual(digitize.default_data_root(Path('C:/x/qls'), {}),
                         Path('C:/x/qls').resolve().parent / 'weconomics-data' / 'olympiads')
